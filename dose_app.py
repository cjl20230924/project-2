"""
空气采样法内照射剂量计算系统 v4.0 (桌面版)
基于 PyQt5 实现，支持四种计算方法：
  1. 国标单AMAD法（默认5μm）
  2. 单峰/双峰拟合法（多模态）
  3. 正态概率图法（直线拟合）
  4. 逐级独立法（每级视为独立均质气溶胶源）

v4.0 新增：
  - 动态分级表格：支持增删分级、自定义粒径范围
  - 逐级化合物配置：每个分级可配置独立的化合物列表
  - 所有计算方法适配动态分级

运行：python dose_app.py
"""

from PyQt5.QtGui import QFont, QColor, QPalette
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QGridLayout, QGroupBox, QLabel, QLineEdit, QComboBox, QPushButton,
    QRadioButton, QButtonGroup, QTableWidget, QTableWidgetItem,
    QSplitter, QScrollArea, QFrame, QDoubleSpinBox,
    QFileDialog, QMessageBox, QTabWidget, QHeaderView, QSizePolicy,
    QStatusBar, QTextEdit
)
from matplotlib.figure import Figure
from matplotlib.backends.backend_qt5agg import FigureCanvasQTAgg as FigureCanvas
import matplotlib.pyplot as plt
import sys, os, warnings, datetime
import numpy as np
import pandas as pd
from pathlib import Path
from scipy.stats import norm
from sklearn.linear_model import LinearRegression
import matplotlib
matplotlib.use('Qt5Agg')
matplotlib.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'DejaVu Sans']
matplotlib.rcParams['axes.unicode_minus'] = False
matplotlib.rcParams['mathtext.fontset'] = 'stix'
warnings.filterwarnings('ignore')


# ==================== 路径 + xf_core 导入 ====================
def resource_path(relative_path):
    if hasattr(sys, '_MEIPASS'):
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), relative_path)

import xf_core
from xf_core import (lognormal_pdf, stage_integral, fit_unimodal, fit_bimodal,
                     fit_linear_probit, recommend_method,
                     calc_unimodal_stats, calc_bimodal_stats, calc_probit_stats,
                     data_quality_assessment)

def _patch_xf_stages(low, high, cut=None):
    old = {'stages_low': xf_core.stages_low.copy(),
           'stages_high': xf_core.stages_high.copy(),
           'cut_diameters': xf_core.cut_diameters.copy(),
           'midpoints': xf_core.midpoints.copy()}
    xf_core.stages_low = np.asarray(low, dtype=float)
    xf_core.stages_high = np.asarray(high, dtype=float)
    xf_core.cut_diameters = np.asarray(cut if cut is not None else low[:-1], dtype=float)
    xf_core.midpoints = np.sqrt(xf_core.stages_low * xf_core.stages_high)
    return old

def _restore_xf_stages(old):
    xf_core.stages_low = old['stages_low']
    xf_core.stages_high = old['stages_high']
    xf_core.cut_diameters = old['cut_diameters']
    xf_core.midpoints = old['midpoints']

# ==================== 默认 9 级 + 常量 ====================
DEFAULT_STAGES = [
    {'name': 'A', 'low': 21.3, 'high': 50.0},
    {'name': 'B', 'low': 14.8, 'high': 21.3},
    {'name': 'C', 'low': 9.8, 'high': 14.8},
    {'name': 'D', 'low': 6.0, 'high': 9.8},
    {'name': 'E', 'low': 3.5, 'high': 6.0},
    {'name': 'F', 'low': 1.6, 'high': 3.5},
    {'name': 'G', 'low': 0.9, 'high': 1.6},
    {'name': 'H', 'low': 0.5, 'high': 0.9},
    {'name': 'Filter', 'low': 0.1, 'high': 0.5},
]
_XF_EFF = np.array([0.52, 0.61, 0.78, 0.89, 0.95, 0.96, 0.97, 0.99, 1.0])
DEFAULT_COMP = {'compound': 'UO2', 'aerosol_type': 'Intermediate Type M/S',
                'activity_fraction': 1.0, 'nuclides': {'U_238': 0.993, 'U_235': 0.007}}
DATA_DIR = Path(resource_path("processed_nuclide_files"))
CONC_COLS = ['A_conc','B_conc','C_conc','D_conc','E_conc','F_conc','G_conc','H_conc','Filter_conc']

MASK_PARAMS = {
    "无防护": {"filtration_efficiency": 0.00, "leakage_rate": 1.00, "is_electret": False, "mpps_um": 0.3},
    "普通熔喷口罩": {"filtration_efficiency": 0.60, "leakage_rate": 0.15, "is_electret": False, "mpps_um": 0.3},
    "KN90": {"filtration_efficiency": 0.90, "leakage_rate": 0.10, "is_electret": True, "mpps_um": 0.3},
    "KN95": {"filtration_efficiency": 0.95, "leakage_rate": 0.08, "is_electret": True, "mpps_um": 0.3},
    "FFP2": {"filtration_efficiency": 0.94, "leakage_rate": 0.08, "is_electret": True, "mpps_um": 0.3},
    "N95": {"filtration_efficiency": 0.95, "leakage_rate": 0.08, "is_electret": True, "mpps_um": 0.3},
    "医用外科口罩": {"filtration_efficiency": 0.70, "leakage_rate": 0.20, "is_electret": True, "mpps_um": 0.3},
}

def _build_efficiency(low, high):
    n = len(low)
    if n == 9:
        d_l = np.array(low); d_h = np.array(high)
        xl = np.array([21.3, 14.8, 9.8, 6.0, 3.5, 1.6, 0.9, 0.5, 0.1])
        xh = np.array([50.0, 21.3, 14.8, 9.8, 6.0, 3.5, 1.6, 0.9, 0.5])
        if np.allclose(d_l, xl, rtol=0.02) and np.allclose(d_h, xh, rtol=0.02):
            return _XF_EFF.copy()
    mids = np.sqrt(np.array(low) * np.array(high))
    xf_mids = np.sqrt(np.array([21.3,14.8,9.8,6.0,3.5,1.6,0.9,0.5,0.1]) *
                      np.array([50,21.3,14.8,9.8,6.0,3.5,1.6,0.9,0.5]))
    return np.interp(mids, xf_mids, _XF_EFF, left=0.85, right=1.0)

def _apply_efficiency_dynamic(low, high, acts):
    return np.asarray(acts, dtype=float) / _build_efficiency(low, high)

# ==================== 核素数据 ====================
_nuclide_cache = {}; _nuclide_files = {}; _element_nuclides_map = {}

def scan_nuclides():
    global _nuclide_files, _element_nuclides_map
    if not DATA_DIR.exists(): return
    for f in list(DATA_DIR.glob("*.parquet")) + list(DATA_DIR.glob("*.xlsx")):
        try:
            df = pd.read_parquet(f) if f.suffix == '.parquet' else pd.read_excel(f)
            cm = {}
            for c in df.columns:
                cl = c.lower().strip()
                if cl == 'element': cm[c] = 'element'
                elif 'radionuclide' in cl: cm[c] = 'radionuclide'
                elif 'route' in cl: cm[c] = 'route_of_intake'
                elif 'aerosol' in cl: cm[c] = 'aerosol_type'
                elif 'compound' in cl: cm[c] = 'compound'
                elif 'particle' in cl: cm[c] = 'particle_size'
                elif 'dose' in cl: cm[c] = 'dose_coefficient'
                elif cl == 'fa': cm[c] = 'fA'
            df = df.rename(columns=cm)
            if 'particle_size' in df.columns:
                df['particle_size'] = pd.to_numeric(df['particle_size'].astype(str)
                    .str.replace(r'\s*[µμ]m\s*|\s*micron\s*', '', regex=True)
                    .replace(['','nan','NaN','None'], np.nan), errors='coerce')
            if 'dose_coefficient' in df.columns: df['dose_coefficient'] = pd.to_numeric(df['dose_coefficient'], errors='coerce')
            if 'aerosol_type' in df.columns: df['aerosol_type'] = df['aerosol_type'].replace('Gaseous', 'Unspecified')
            if 'element' in df.columns and 'radionuclide' in df.columns:
                elem = str(df['element'].iloc[0]); nuc = str(df['radionuclide'].iloc[0]).replace('_','-')
                _nuclide_files[nuc] = (f, df)
                _element_nuclides_map.setdefault(elem, []); _element_nuclides_map[elem].append(nuc)
        except Exception as e: print(f"[警告] 加载 {f.name} 失败: {e}")

def _normalize_nuclide(name): return name.strip().replace('_', '-')

_COMPOUND_SHORT_TO_FULL = {
    'UO2': ('Uranium octoxide, uranium dioxide', 'Intermediate Type M/S'),
    'U3O8': ('Uranium octoxide, uranium dioxide', 'Intermediate Type M/S'),
    'UO3': ('Uranyl nitrate, uranium peroxide hydrate, ammonium diuranate, uranium trioxide', 'Intermediate Type F/M'),
    '硝酸铀酰': ('Uranyl nitrate, uranium peroxide hydrate, ammonium diuranate, uranium trioxide', 'Intermediate Type F/M'),
    'UNH': ('Uranyl nitrate, uranium peroxide hydrate, ammonium diuranate, uranium trioxide', 'Intermediate Type F/M'),
    'ADU': ('Uranyl nitrate, uranium peroxide hydrate, ammonium diuranate, uranium trioxide', 'Intermediate Type F/M'),
    'UF6': ('Uranium hexafluoride, uranyl tributyl-phosphate', 'Type F'),
    'U_metal': ('Uranyl acetylacetonate; depleted uranium aerosols...', 'Type M'),
    'DU': ('Uranyl acetylacetonate; depleted uranium aerosols...', 'Type M'),
    'PuO2': ('Plutonium-239 dioxide, plutonium in mixed oxide', 'Unspecified'),
    'MOX': ('Plutonium-239 dioxide, plutonium in mixed oxide', 'Unspecified'),
    'Pu_nitrate': ('Plutonium nitrate', 'Unspecified'),
    'Pu_citrate': ('Plutonium citrate, plutonium tri-butyl-phosphate, plutonium chloride', 'Type M'),
    'HTO': ('Gas or vapour Type V, Tritiated water', 'Unspecified'),
    'HT': ('Gas or vapour Type V, Tritium gas', 'Unspecified'),
    'OBT': ('Biogenic organic compounds', 'Unspecified'),
}

def resolve_compound(short):
    r = _COMPOUND_SHORT_TO_FULL.get(short.strip())
    return r if r else (short.strip(), None)

def get_nuclide_df(nuclide, inhalation_only=True):
    nk = _normalize_nuclide(nuclide); key = (nk, inhalation_only)
    if key in _nuclide_cache: return _nuclide_cache[key]
    if nk not in _nuclide_files: return None
    _, df = _nuclide_files[nk]
    if inhalation_only and 'route_of_intake' in df.columns:
        df = df[df['route_of_intake'] == 'Inhalation'].copy()
    _nuclide_cache[key] = df; return df

def interp_dose_coeff(df, aero, ps_um):
    if df is None or df.empty: return None
    if 'aerosol_type' not in df.columns: return None
    sub = df[df['aerosol_type'] == aero].copy()
    if sub.empty: return None
    sub = sub.dropna(subset=['particle_size', 'dose_coefficient']).sort_values('particle_size')
    if sub.empty:
        r0 = df[df['aerosol_type'] == aero]
        return r0['dose_coefficient'].iloc[0] if not r0.empty else None
    ps = sub['particle_size'].values; dcs = sub['dose_coefficient'].values
    if len(ps) == 1: return dcs[0]
    if ps_um <= ps.min(): return dcs[0]
    if ps_um >= ps.max(): return dcs[-1]
    lx = np.log(ps_um); lp = np.log(ps)
    for i in range(len(lp)-1):
        if lp[i] <= lx <= lp[i+1]:
            y1, y2 = dcs[i], dcs[i+1]; x1, x2 = lp[i], lp[i+1]
            return y1 + (lx-x1)/(x2-x1)*(y2-y1) if x2 != x1 else y1
    return None

def calc_mask_pf(mt, ps_um=None):
    if mt == "无防护": return 1.0
    p = MASK_PARAMS.get(mt, MASK_PARAMS["无防护"])
    be, lk, mp, ie = p["filtration_efficiency"], p["leakage_rate"], p["mpps_um"], p["is_electret"]
    if ps_um is None: ae = be
    else:
        sz = ps_um
        if not ie: ae = be * min(1.0, sz/mp) if sz > 0.01 else 0.0
        else:
            if sz <= 0.01: ae = 0.1
            elif abs(sz-mp) < 1e-6: ae = be
            else: ae = min(be + min(abs(sz-mp)/mp*0.15, 0.09), 0.99)
    return lk + (1-ae) * (1-lk)

# ==================== CalcThread ====================
class CalcThread(QThread):
    result_ready = pyqtSignal(dict); error_signal = pyqtSignal(str)
    def __init__(self, func, *a, **kw): super().__init__(); self.func = func; self.args = a; self.kwargs = kw
    def run(self):
        try: self.result_ready.emit(self.func(*self.args, **self.kwargs))
        except Exception as e:
            import traceback; self.error_signal.emit(traceback.format_exc())

# ==================== PlotCanvas（支持动态分级）====================
class PlotCanvas(FigureCanvas):
    def __init__(self, parent=None, w=7, h=7):
        self.fig = Figure(figsize=(w, h), dpi=88)
        super().__init__(self.fig); self.setParent(parent)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)

    def plot_fitting(self, raw, corr, a_u, g_u, t_u, a1, g1, f1, a2, g2, tb, tc,
                     D50, GSD, R2, is_corr=False, dl=None, dh=None, dm=None, dc=None):
        _l = dl if dl is not None else xf_core.stages_low
        _h = dh if dh is not None else xf_core.stages_high
        _m = dm if dm is not None else xf_core.midpoints
        _c = dc if dc is not None else xf_core.cut_diameters
        self.fig.clear()
        dlbl = "防护后校正浓度" if is_corr else "防护前浓度（效率校正）"
        pn = "（图中散点 = 口罩防护校正后数据）" if is_corr else "（图中散点 = 防护前数据，效率校正）"
        xp = np.logspace(np.log10(0.1), np.log10(50), 300)
        sw = np.log(_h/_l)
        dd = corr/sw/tc if tc > 0 else corr
        yu = lognormal_pdf(xp, a_u, g_u)*t_u/tc if tc > 0 else np.zeros_like(xp)
        axes = self.fig.subplots(2, 2); sc = '#e74c3c' if is_corr else '#3498db'
        # 左上：单峰
        ax = axes[0,0]; ax.semilogx(xp, yu, 'b-', lw=2, label=f'单峰 AMAD={a_u:.2f} $\mu m$')
        ax.scatter(_m, dd, c=sc, s=50, edgecolors='k', zorder=5, label=dlbl)
        ax.set_xlabel('空气动力学粒径 ($\mu m$)'); ax.set_ylabel('归一化密度')
        ax.set_title(f'单峰对数正态拟合  [{dlbl}]'); ax.grid(ls='--', alpha=0.5); ax.set_xlim(0.1, 50); ax.set_ylim(bottom=0)
        ax.legend(fontsize=7, loc='upper left')
        # 右上：双峰
        ax = axes[0,1]
        if not np.isnan(a2):
            yc = f1 * lognormal_pdf(xp, a1, g1); yf = (1-f1) * lognormal_pdf(xp, a2, g2)
            ax.semilogx(xp, yc+yf, 'k-', lw=2, label='总拟合')
            ax.semilogx(xp, yc, 'b--', lw=1.5, label=f'粗峰 {a1:.1f} $\mu m$')
            ax.semilogx(xp, yf, 'r--', lw=1.5, label=f'细峰 {a2:.1f} $\mu m$')
        else: ax.semilogx(xp, yu, 'k--', lw=2, label='未检测到双峰')
        ax.scatter(_m, dd, c=sc, s=50, edgecolors='k', zorder=5, label=dlbl)
        ax.set_xlabel('空气动力学粒径 ($\mu m$)'); ax.set_title(f'双峰对数正态拟合  [{dlbl}]')
        ax.grid(ls='--', alpha=0.5); ax.set_xlim(0.1, 50); ax.set_ylim(bottom=0)
        ax.legend(fontsize=7, loc='upper left')
        # 左下：正态概率图
        ax = axes[1,0]
        ra = raw[:-1] if len(raw) > 1 else raw; fa = raw[-1] if len(raw) > 0 else 0
        tr = np.sum(ra)+fa
        if tr > 0 and len(_c) >= 2:
            f = ra/tr; fall = np.concatenate([f, [fa/tr]])
            cff = np.cumsum(np.flip(fall))[:-1]; cl = np.flip(cff)*100
            vld = (cl>1)&(cl<99)
            if np.sum(vld) >= 3:
                xv = np.log(_c[vld]); yv = norm.ppf(cl[vld]/100)
                reg = LinearRegression().fit(xv.reshape(-1,1), yv)
                sl, ic = reg.coef_[0], reg.intercept_
                la = -ic/sl; AP = np.exp(la); lg = 1/sl; GP = np.exp(lg)
                D84 = np.exp(la+lg); D159 = np.exp(la-lg)
                ax.set_xscale('log'); ax.set_xlim(0.3, 30)
                pts = [1,5,10,20,30,40,50,60,70,80,90,95,99]
                ptt = norm.ppf(np.array(pts)/100); ax.set_yticks(ptt)
                ax.set_yticklabels([str(p) for p in pts], fontsize=7)
                ax.set_ylim(norm.ppf(0.005), norm.ppf(0.995))
                ax.scatter(_c[vld], yv, c='red', s=60, label='实测数据', zorder=5)
                xf2 = np.logspace(np.log10(0.5), np.log10(21.3), 200)
                ax.plot(xf2, sl*np.log(xf2)+ic, 'b-', lw=2, label='ICRP 回归线')
                ax.scatter(AP, 0, c='green', s=90, marker='s', label=f'$AMAD$={AP:.2f} $\mu m$', zorder=6)
                ax.scatter(D84, norm.ppf(0.8413), c='orange', s=90, marker='s', label=f'$D_{{84.1}}$={D84:.1f} $\mu m$', zorder=6)
                ax.scatter(D159, norm.ppf(0.1587), c='purple', s=90, marker='s', label=f'$D_{{15.9}}$={D159:.1f} $\mu m$', zorder=6)
                ax.axhline(y=0, color='gray', ls=':', alpha=0.5)
                sr = np.sum((yv-reg.predict(xv.reshape(-1,1)))**2); st = np.sum((yv-np.mean(yv))**2)
                r2c = 1-sr/st if st > 0 else 0
                ax.text(0.05, 0.95, f'$AMAD$={AP:.2f} $\mu m$\n$GSD$={GP:.2f}\n$R^2$={r2c:.4f}',
                        transform=ax.transAxes, va='top', fontsize=8,
                        bbox=dict(boxstyle='round', fc='wheat', alpha=0.7))
        ax.set_xlabel('空气动力学粒径 ($\mu m$)'); ax.set_ylabel('累积活度比例 (%)')
        ax.set_title(f'正态概率图（ICRP方法）  [{dlbl}]'); ax.grid(True, which='both', ls='--', alpha=0.5)
        ax.legend(fontsize=6.5, loc='lower right')
        # 右下：摘要
        ax = axes[1,1]; ax.axis('off')
        ls = ['══════ 拟合参数摘要 ══════', f'数据类型: {dlbl}', '', '【单峰对数正态】',
              f'  $AMAD$ = {a_u:.3f} $\\mu$m', f'  $GSD$  = {g_u:.3f}', '']
        if not np.isnan(a2):
            ls += ['【双峰对数正态】', f'  粗峰 $AMAD$ = {a1:.3f}', f'  粗峰 $GSD$  = {g1:.3f}',
                   f'  细峰 $AMAD$ = {a2:.3f}', f'  细峰 $GSD$  = {g2:.3f}', f'  粗峰占比  = {f1:.2%}']
        else: ls += ['【双峰对数正态】', '  (未检测到)']
        ls.append('')
        if not np.isnan(D50) and not np.isnan(GSD):
            ls += ['【正态概率图】', f'  $AMAD$ = {D50:.3f}', f'  $GSD$  = {GSD:.3f}', f'  $R^2$   = {R2:.4f}']
        else: ls += ['【正态概率图】', '  (数据点不足)']
        ax.text(0.1, 0.95, '\n'.join(ls), transform=ax.transAxes, ha='left', va='top', fontsize=8.5,
                bbox=dict(boxstyle='round', fc='#f0f4ff', alpha=0.9, ec='#6688aa'))
        self.fig.tight_layout(pad=2.0); self.draw()

# ==================== 紧凑化合物卡片 ====================
class CompactCompoundCard(QFrame):
    """每级可配置多个化合物；核素从数据文件下拉选取（不再手动键入）"""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setFrameShape(QFrame.StyledPanel)
        self.setStyleSheet("QFrame { border:1px solid #cce0ff; border-radius:4px; background:#f8fbff; }")
        o = QVBoxLayout(self); o.setContentsMargins(4,2,4,2); o.setSpacing(2)
        # Row 1: 化合物 + 占比
        r1 = QHBoxLayout(); r1.setSpacing(4)
        r1.addWidget(QLabel("化合物:")); self.ce = QLineEdit("UO2"); self.ce.setMinimumWidth(70); r1.addWidget(self.ce, 1)
        r1.addWidget(QLabel("占比:")); self.af = QDoubleSpinBox()
        self.af.setRange(0,1); self.af.setSingleStep(0.05); self.af.setValue(1.0); self.af.setDecimals(3); self.af.setMinimumWidth(60)
        r1.addWidget(self.af, 1); o.addLayout(r1)
        # Row 2: 气溶胶类型
        r2 = QHBoxLayout(); r2.setSpacing(4)
        r2.addWidget(QLabel("气溶胶:")); self.ac = QComboBox(); self.ac.setMinimumWidth(100); r2.addWidget(self.ac, 1); o.addLayout(r2)
        # Row 3: 核素下拉 + 丰度 + 添加按钮
        r3 = QHBoxLayout(); r3.setSpacing(4)
        r3.addWidget(QLabel("核素:")); self.nc = QComboBox(); self.nc.setMinimumWidth(85)
        self.nc.setToolTip("从已加载的核素数据文件中选取"); r3.addWidget(self.nc, 1)
        r3.addWidget(QLabel("丰度:")); self.nab = QDoubleSpinBox()
        self.nab.setRange(0.001, 1.0); self.nab.setSingleStep(0.01); self.nab.setValue(1.0); self.nab.setDecimals(4); self.nab.setMinimumWidth(55)
        r3.addWidget(self.nab, 1)
        self.nadd = QPushButton("＋"); self.nadd.setFixedSize(24, 22)
        self.nadd.setStyleSheet("QPushButton { font-size:11px; font-weight:bold; color:#27ae60; border:1px solid #a0d8b0; border-radius:3px; background:#e8f8ee; } QPushButton:hover { background:#d0f0d8; }")
        r3.addWidget(self.nadd); o.addLayout(r3)
        # 核素已选列表区
        self._nlw = QWidget(); self._nll = QVBoxLayout(self._nlw); self._nll.setSpacing(1); self._nll.setContentsMargins(22,0,0,0)
        o.addWidget(self._nlw)
        self._nuclides = {}          # {name: abundance}
        self._nrows = []             # [(row_widget, label, del_btn), …]
        self._on_changed_cb = None   # 外部回调（用于通知父级数据变更）
        self.nadd.clicked.connect(self._on_add_nuclide)

    # ── 核素增删 ──
    def _on_add_nuclide(self):
        name = self.nc.currentText()
        if not name: return
        ab = self.nab.value()
        self._nuclides[name] = ab
        self._rebuild_nuclide_list()
        if self._on_changed_cb: self._on_changed_cb()

    def _on_del_nuclide(self, name):
        if name in self._nuclides:
            del self._nuclides[name]
            self._rebuild_nuclide_list()
            if self._on_changed_cb: self._on_changed_cb()

    def _rebuild_nuclide_list(self):
        for rw, _, _ in self._nrows:
            self._nll.removeWidget(rw); rw.deleteLater()
        self._nrows.clear()
        for nm, ab in self._nuclides.items():
            rw = QWidget(); rl = QHBoxLayout(rw); rl.setContentsMargins(0,0,0,0); rl.setSpacing(3)
            lb = QLabel(f"{nm}: {ab:.4f}"); lb.setStyleSheet("font-size:10px; color:#1a5276;")
            db = QPushButton("✕"); db.setFixedSize(18, 18)
            db.setStyleSheet("QPushButton { font-size:9px; color:#c0392b; border:1px solid #e8c0c0; border-radius:2px; background:#fff5f5; } QPushButton:hover { background:#ffe0e0; }")
            db.clicked.connect(lambda checked, n=nm: self._on_del_nuclide(n))
            rl.addWidget(lb, 1); rl.addWidget(db)
            self._nll.addWidget(rw); self._nrows.append((rw, lb, db))

    # ── 数据接口 ──
    def get_data(self):
        c = self.ce.text().strip()
        r = _COMPOUND_SHORT_TO_FULL.get(c)
        aero = r[1] if r else self.ac.currentText()
        return {'compound': c, 'aerosol_type': aero,
                'activity_fraction': self.af.value(),
                'nuclides': dict(self._nuclides)}

    def set_data(self, d):
        self.ce.setText(d.get('compound', 'UO2'))
        self.af.setValue(d.get('activity_fraction', 1.0))
        at = d.get('aerosol_type', 'Intermediate Type M/S'); i = self.ac.findText(at)
        if i >= 0: self.ac.setCurrentIndex(i)
        self._nuclides = dict(d.get('nuclides', {'U_238': 0.993, 'U_235': 0.007}))
        self._rebuild_nuclide_list()

    def update_aerosol_options(self, opts):
        cur = self.ac.currentText(); self.ac.clear(); self.ac.addItems(opts)
        i = self.ac.findText(cur)
        if i >= 0: self.ac.setCurrentIndex(i)

    def update_nuclide_options(self, nuclides):
        """由外部（DoseCalcApp）根据所选元素注入可用的核素列表"""
        cur = self.nc.currentText()
        self.nc.blockSignals(True); self.nc.clear()
        if nuclides: self.nc.addItems(sorted(nuclides))
        self.nc.blockSignals(False)
        if cur and self.nc.findText(cur) >= 0:
            self.nc.setCurrentText(cur)

# ==================== 主窗口 ====================
class DoseCalcApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("空气采样法内照射剂量计算系统 v4.0")
        self.setMinimumSize(1100, 700); self.resize(1440, 860)
        self._fit_result = None; self._calc_history = []; self._calc_thread = None; self._file_df = None
        self._stages = []; self._init_default_stages()
        self._selected_stage_idx = 0; self._compound_cards = []; self._current_elem = ''
        self._setup_ui(); self._scan_data()

    def _init_default_stages(self):
        self._stages = []
        for ds in DEFAULT_STAGES:
            self._stages.append({**ds, 'conc': 0.0, 'compounds': [dict(DEFAULT_COMP)]})

    def _get_stage_arrays(self):
        l = np.array([s['low'] for s in self._stages], dtype=float)
        h = np.array([s['high'] for s in self._stages], dtype=float)
        m = np.sqrt(l*h); n = len(l)
        c = h[:-1].copy() if n > 1 else np.array([])
        concs = np.array([s['conc'] for s in self._stages], dtype=float)
        return l, h, m, c, concs

    # ── UI ──
    def _setup_ui(self):
        cw = QWidget(); self.setCentralWidget(cw); rl = QVBoxLayout(cw); rl.setSpacing(4); rl.setContentsMargins(8,4,8,4)
        tt = QLabel("空气采样法 · 内照射剂量计算系统  v4.0")
        tt.setAlignment(Qt.AlignCenter); tt.setFont(QFont("微软雅黑", 14, QFont.Bold))
        tt.setStyleSheet("color:#1a5276; padding:3px 0;"); rl.addWidget(tt)
        self._splitter = QSplitter(Qt.Horizontal); rl.addWidget(self._splitter, 1)
        # 左侧
        ls = QScrollArea(); ls.setWidgetResizable(True); ls.setMinimumWidth(400)
        lc = QWidget(); ll = QVBoxLayout(lc); ll.setSpacing(7); ll.setContentsMargins(6,6,6,6); ls.setWidget(lc)
        self._splitter.addWidget(ls)
        # 右侧
        rp = QWidget(); rrl = QVBoxLayout(rp); rrl.setSpacing(4); rrl.setContentsMargins(4,0,4,0); self._splitter.addWidget(rp)
        self._splitter.setSizes([600,990]); self._splitter.setStretchFactor(0,0); self._splitter.setStretchFactor(1,1)
        self._splitter.setCollapsible(0,False); self._splitter.setCollapsible(1,False); self._splitter.setHandleWidth(6)
        self._splitter.setStyleSheet("QSplitter::handle { background:#c8d8ea; border-radius:3px; } QSplitter::handle:hover { background:#3498db; }")

        # ─ § 1. 采样数据输入（动态分级 + 化合物）─
        g1 = QGroupBox("① 采样数据输入"); g1.setFont(QFont("微软雅黑", 10, QFont.Bold)); gg1 = QVBoxLayout(g1); gg1.setSpacing(5)
        sr = QHBoxLayout()
        self.rb_manual = QRadioButton("手动输入"); self.rb_file = QRadioButton("从 CSV / Excel 读取"); self.rb_manual.setChecked(True)
        bg = QButtonGroup(self); bg.addButton(self.rb_manual); bg.addButton(self.rb_file)
        sr.addWidget(self.rb_manual); sr.addWidget(self.rb_file); sr.addStretch(); gg1.addLayout(sr)
        # 文件行
        self.fr = QWidget(); fg = QGridLayout(self.fr); fg.setContentsMargins(0,0,0,0); fg.setSpacing(4)
        self.fp_edit = QLineEdit(); self.fp_edit.setPlaceholderText("选择 CSV / Excel 文件…")
        bb = QPushButton("浏览…"); bb.setFixedWidth(56); bb.clicked.connect(self._browse_file)
        self.ws_cb = QComboBox(); self.ws_cb.setMinimumWidth(90)
        self.sid_cb = QComboBox(); self.sid_cb.setMinimumWidth(90)
        self.ws_cb.currentTextChanged.connect(self._on_ws_changed)
        self.sid_cb.currentTextChanged.connect(self._on_sid_changed)
        fg.addWidget(QLabel("文件:"),0,0); fg.addWidget(self.fp_edit,0,1,1,3); fg.addWidget(bb,0,4)
        fg.addWidget(QLabel("车间:"),1,0); fg.addWidget(self.ws_cb,1,1)
        fg.addWidget(QLabel("采样ID:"),1,2); fg.addWidget(self.sid_cb,1,3,1,2)
        self.fr.setVisible(False); gg1.addWidget(self.fr)
        # 分级表格
        gg1.addWidget(QLabel("分级配置（粒径范围 + 活度浓度）："))
        self.st = QTableWidget(0, 5)
        self.st.setHorizontalHeaderLabels(["序号","粒径下限(μm)","粒径上限(μm)","活度浓度(Bq/m³)","化合物"])
        self.st.verticalHeader().setVisible(False); self.st.setSelectionBehavior(QTableWidget.SelectRows)
        self.st.setSelectionMode(QTableWidget.SingleSelection); self.st.setMinimumHeight(180); self.st.setMaximumHeight(260)
        self.st.setAlternatingRowColors(True)
        self.st.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        self.st.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.st.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.st.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        self.st.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeToContents)
        self.st.setColumnWidth(3, 140)
        self.st.itemSelectionChanged.connect(self._on_stage_selection_changed); gg1.addWidget(self.st)
        # 按钮行
        sbr = QHBoxLayout()
        self.btn_add_s = QPushButton("＋ 添加分级"); self.btn_del_s = QPushButton("－ 删除选中分级")
        self.btn_cp = QPushButton("📋 复制化合物到所有分级")
        self.btn_add_s.clicked.connect(self._add_stage); self.btn_del_s.clicked.connect(self._del_stage)
        self.btn_cp.clicked.connect(self._copy_compounds_to_all)
        sbr.addWidget(self.btn_add_s); sbr.addWidget(self.btn_del_s); sbr.addWidget(self.btn_cp); sbr.addStretch()
        gg1.addLayout(sbr)
        # 化合物区
        gg1.addWidget(QLabel("当前选中分级的化合物配置："))
        self._sel_lbl = QLabel("未选中任何分级"); self._sel_lbl.setStyleSheet("color:#888; font-size:10px;"); gg1.addWidget(self._sel_lbl)
        ecr = QHBoxLayout(); ecr.addWidget(QLabel("元素:"))
        self.ec = QComboBox(); self.ec.setMinimumWidth(80); self.ec.currentTextChanged.connect(self._on_elem_changed)
        ecr.addWidget(self.ec); ecr.addStretch()
        bac = QPushButton("＋ 添加化合物到当前分级"); bac.setFixedWidth(180); bac.clicked.connect(self._add_compound_to_current_stage)
        ecr.addWidget(bac); gg1.addLayout(ecr)
        cs = QScrollArea(); cs.setWidgetResizable(True); cs.setFrameShape(QFrame.NoFrame)
        cs.setMinimumHeight(100); cs.setMaximumHeight(200)
        self._cc = QWidget(); self._cl = QVBoxLayout(self._cc); self._cl.setSpacing(3); self._cl.setContentsMargins(0,0,0,0)
        self._cl.addStretch(); cs.setWidget(self._cc); gg1.addWidget(cs)
        self.rb_manual.toggled.connect(self._on_input_mode_changed); self.rb_file.toggled.connect(self._on_input_mode_changed)
        ll.addWidget(g1)
        # ─ § 2. 防护参数 ─
        g4 = QGroupBox("② 防护 & 呼吸参数"); g4.setFont(QFont("微软雅黑", 10, QFont.Bold)); g4l = QGridLayout(g4); g4l.setSpacing(6)
        g4l.addWidget(QLabel("口罩型号:"),0,0)
        self.mc = QComboBox(); self.mc.addItems(list(MASK_PARAMS.keys())); g4l.addWidget(self.mc,0,1)
        self.mil = QLabel(""); self.mil.setStyleSheet("color:#555; font-size:10px;"); g4l.addWidget(self.mil,0,2,1,2)
        self.mc.currentTextChanged.connect(self._update_mask_info)
        g4l.addWidget(QLabel("呼吸速率 BR (m³/h):"),1,0)
        self.brs = QDoubleSpinBox(); self.brs.setRange(0.1,5); self.brs.setValue(1.2); self.brs.setSingleStep(0.1); self.brs.setDecimals(2); g4l.addWidget(self.brs,1,1)
        g4l.addWidget(QLabel("工作时长 T (h):"),1,2)
        self.ws_spin = QDoubleSpinBox(); self.ws_spin.setRange(0.5,24); self.ws_spin.setValue(8.0); self.ws_spin.setSingleStep(0.5); self.ws_spin.setDecimals(1); g4l.addWidget(self.ws_spin,1,3)
        self._update_mask_info(); ll.addWidget(g4)
        # 拟合按钮
        fbr = QHBoxLayout(); fbr.addStretch()
        self.fb = QPushButton("🔍  运行拟合（查看粒径分布图）"); self.fb.setMinimumHeight(36)
        self.fb.setFont(QFont("微软雅黑", 11, QFont.Bold))
        self.fb.setStyleSheet("QPushButton { background:#27ae60; color:white; border-radius:5px; padding:4px 20px; } QPushButton:hover { background:#2ecc71; } QPushButton:disabled { background:#aaa; }")
        self.fb.clicked.connect(self._on_fit); fbr.addWidget(self.fb)
        fh = QLabel("  ↳ 先运行拟合可在右侧查看粒径分布图，再选择计算方法开始计算")
        fh.setStyleSheet("color:#888; font-size:10px;"); ll.addLayout(fbr); ll.addWidget(fh)
        # ─ § 3. 计算方法 ─
        g2 = QGroupBox("③ 计算方法选择"); g2.setFont(QFont("微软雅黑", 10, QFont.Bold)); g2l = QVBoxLayout(g2); g2l.setSpacing(3)
        self.rb_std = QRadioButton("方法1 · 国标单AMAD（固定 5 μm）")
        self.rb_mod = QRadioButton("方法2 · 多模态拟合（手动选择峰型）")
        self.rb_pro = QRadioButton("方法3 · 正态概率图法（直线拟合）")
        self.rb_stg = QRadioButton("方法4 · 逐级独立法（每级视为均质源）")
        self.rb_std.setChecked(True); bg2 = QButtonGroup(self)
        _h = {self.rb_std:"  ↳ 使用固定 AMAD = 5 μm，符合 ICRP‑66 及国标。",
              self.rb_mod:"  ↳ 对分级数据拟合对数正态分布，由您在下方选择单峰或双峰。",
              self.rb_pro:"  ↳ 用累积活度–正态概率图进行直线回归，求 AMAD 与 GSD。",
              self.rb_stg:"  ↳ 每级视为独立气溶胶源，利用该级中值粒径直接查表后逐级累加。"}
        self._hl = {}
        self._mpw = QWidget(); mpr = QHBoxLayout(self._mpw); mpr.setContentsMargins(22,2,0,2); mpr.setSpacing(14)
        mpr.addWidget(QLabel("峰型选择：")); self.rb_pu = QRadioButton("单峰（Unimodal）"); self.rb_pb = QRadioButton("双峰（Bimodal）")
        self.rb_pu.setChecked(True); bgp = QButtonGroup(self); bgp.addButton(self.rb_pu); bgp.addButton(self.rb_pb)
        mpr.addWidget(self.rb_pu); mpr.addWidget(self.rb_pb)
        mpr.addWidget(QLabel("（拟合图运行后可在右侧查看双峰结果，再决定选哪种）")); mpr.addStretch()
        self._mpw.setVisible(False)
        for rb, hint in _h.items():
            bg2.addButton(rb); g2l.addWidget(rb); hl = QLabel(hint)
            hl.setStyleSheet("color:#888; font-size:10px; margin-left:16px;"); hl.setWordWrap(True)
            g2l.addWidget(hl); self._hl[rb] = hl
            if rb is self.rb_mod: g2l.addWidget(self._mpw)
            rb.toggled.connect(self._update_method_hints)
        self._update_method_hints(); ll.addWidget(g2)
        # 计算按钮
        btr = QHBoxLayout()
        self.cb = QPushButton("▶  开始计算"); self.cb.setMinimumHeight(40)
        self.cb.setFont(QFont("微软雅黑", 12, QFont.Bold))
        self.cb.setStyleSheet("QPushButton { background:#1a5276; color:white; border-radius:6px; } QPushButton:hover { background:#2471a3; } QPushButton:disabled { background:#aaa; }")
        self.cb.clicked.connect(self._on_calc)
        bc = QPushButton("清空结果"); bc.setMinimumHeight(40); bc.setFixedWidth(90); bc.clicked.connect(self._clear_results)
        btr.addWidget(self.cb, 1); btr.addWidget(bc); ll.addLayout(btr); ll.addStretch()

        # ═══════ 右侧 Tab ═══════
        rt = QTabWidget(); rt.setDocumentMode(True); rrl.addWidget(rt, 1)
        # Tab1: 拟合图
        tf = QWidget(); tfl = QVBoxLayout(tf); tfl.setContentsMargins(4,4,4,4)
        fs = QSplitter(Qt.Vertical); fs.setHandleWidth(6)
        fs.setStyleSheet("QSplitter::handle { background:#c8d8ea; border-radius:3px; } QSplitter::handle:hover { background:#3498db; }")
        self.pc = PlotCanvas(tf, 7, 7)
        self.fil = QLabel("拟合参数将在计算后显示")
        self.fil.setStyleSheet("background:#eaf2ff; border:1px solid #aac; padding:6px; border-radius:4px;"); self.fil.setWordWrap(True); self.fil.setMinimumHeight(40)
        fs.addWidget(self.pc); fs.addWidget(self.fil); fs.setSizes([560,80])
        fs.setCollapsible(0,False); fs.setCollapsible(1,False); tfl.addWidget(fs, 1); rt.addTab(tf, "📈 拟合图")
        # Tab2: 剂量结果
        td_w = QWidget(); tdl = QVBoxLayout(td_w); tdl.setContentsMargins(4,4,4,4); tdl.setSpacing(4)
        ds = QSplitter(Qt.Vertical); ds.setHandleWidth(6)
        ds.setStyleSheet("QSplitter::handle { background:#c8d8ea; border-radius:3px; } QSplitter::handle:hover { background:#3498db; }")
        topw = QWidget(); tl = QVBoxLayout(topw); tl.setContentsMargins(0,0,0,0); tl.setSpacing(4)
        self.sl = QLabel("剂量结果将在计算后显示")
        self.sl.setStyleSheet("background:#eafaf1; border:1px solid #aad; padding:8px; font-size:13px; border-radius:4px;"); self.sl.setWordWrap(True)
        tl.addWidget(self.sl)
        self.rt2 = QTableWidget(0,7); self.rt2.setHorizontalHeaderLabels(["化合物","气溶胶类型","核素","AMAD/粒径","e (Sv/Bq)","活度贡献","剂量 (Sv)"])
        self.rt2.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.rt2.setAlternatingRowColors(True)
        self.rt2.setEditTriggers(QTableWidget.NoEditTriggers); tl.addWidget(self.rt2, 1); ds.addWidget(topw)
        botw = QWidget(); bl = QVBoxLayout(botw); bl.setContentsMargins(0,0,0,0); bl.setSpacing(4)
        bl.addWidget(QLabel("累计剂量列表（可多次计算叠加）"))
        self.atl = QLabel(""); self.atl.setStyleSheet("color:#1a5276; font-weight:bold;"); bl.addWidget(self.atl)
        self.at = QTableWidget(0,5); self.at.setHorizontalHeaderLabels(["方法","化合物/核素","AMAD (μm)","口罩防护因子","剂量 (Sv)"])
        self.at.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.at.setAlternatingRowColors(True); self.at.setEditTriggers(QTableWidget.NoEditTriggers); bl.addWidget(self.at, 1)
        abr = QHBoxLayout()
        ba = QPushButton("➕ 将本次结果加入累计"); bd = QPushButton("➖ 删除选中行"); bcl = QPushButton("🗑 清空累计")
        ba.clicked.connect(self._add_to_accum); bd.clicked.connect(self._del_accum_row); bcl.clicked.connect(self._clear_accum)
        abr.addWidget(ba); abr.addWidget(bd); abr.addWidget(bcl); abr.addStretch(); bl.addLayout(abr)
        ge = QGroupBox("④ 导出结果"); ge.setFont(QFont("微软雅黑", 10, QFont.Bold)); el = QHBoxLayout(ge); el.setSpacing(8)
        self.bx = QPushButton("📥 Excel"); self.bx.setMinimumHeight(34); self.bx.clicked.connect(self._export_xlsx)
        self.bc2 = QPushButton("📄 CSV"); self.bc2.setMinimumHeight(34); self.bc2.clicked.connect(self._export_csv)
        self.bl2 = QPushButton("📝 日志"); self.bl2.setMinimumHeight(34); self.bl2.clicked.connect(self._export_log)
        self.ba2 = QPushButton("📦 全部（含图）"); self.ba2.setMinimumHeight(34); self.ba2.clicked.connect(self._export_all)
        el.addWidget(self.bx); el.addWidget(self.bc2); el.addWidget(self.bl2); el.addWidget(self.ba2); el.addStretch()
        bl.addWidget(ge); ds.addWidget(botw); ds.setSizes([320,250])
        ds.setCollapsible(0,False); ds.setCollapsible(1,False); tdl.addWidget(ds, 1); rt.addTab(td_w, "💊 剂量结果")
        # Tab3: 日志
        self.le = QTextEdit(); self.le.setReadOnly(True); self.le.setFont(QFont("Consolas", 9)); rt.addTab(self.le, "📋 计算日志")
        self.sb = QStatusBar(); self.setStatusBar(self.sb); self.sb.showMessage("就绪")
        self._rebuild_stage_table()

    # ── 分级表格操作 ──
    def _rebuild_stage_table(self):
        self.st.blockSignals(True); self.st.setRowCount(len(self._stages))
        for i, s in enumerate(self._stages):
            it0 = QTableWidgetItem(s['name']); it0.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable); it0.setTextAlignment(Qt.AlignCenter)
            self.st.setItem(i, 0, it0)
            sl = QDoubleSpinBox(); sl.setRange(0.001, 100); sl.setDecimals(2); sl.setValue(s['low'])
            sl.setButtonSymbols(QDoubleSpinBox.NoButtons)
            sl.valueChanged.connect(lambda v, idx=i: self._on_sp(idx, 'low', v)); self.st.setCellWidget(i, 1, sl)
            sh = QDoubleSpinBox(); sh.setRange(0.001, 100); sh.setDecimals(2); sh.setValue(s['high'])
            sh.setButtonSymbols(QDoubleSpinBox.NoButtons)
            sh.valueChanged.connect(lambda v, idx=i: self._on_sp(idx, 'high', v)); self.st.setCellWidget(i, 2, sh)
            sc = QDoubleSpinBox(); sc.setRange(0, 1e9); sc.setDecimals(4); sc.setValue(s['conc'])
            sc.setButtonSymbols(QDoubleSpinBox.NoButtons)
            sc.valueChanged.connect(lambda v, idx=i: self._on_sp(idx, 'conc', v)); self.st.setCellWidget(i, 3, sc)
            nc = len(s.get('compounds', [])); itc = QTableWidgetItem(str(nc))
            itc.setFlags(Qt.ItemIsEnabled | Qt.ItemIsSelectable); itc.setTextAlignment(Qt.AlignCenter); self.st.setItem(i, 4, itc)
            self.st.setRowHeight(i, 30)
        self.st.blockSignals(False); self._update_compound_ui()

    def _on_sp(self, idx, key, val):
        if 0 <= idx < len(self._stages): self._stages[idx][key] = val

    def _add_stage(self):
        n = len(self._stages); self._stages.append({'name': f"Stage{n+1}", 'low': 1.0, 'high': 3.0, 'conc': 0.0, 'compounds': [dict(DEFAULT_COMP)]})
        self._rebuild_stage_table(); self.st.selectRow(len(self._stages)-1)

    def _del_stage(self):
        sel = self.st.selectedItems()
        if not sel: return QMessageBox.information(self, "提示", "请先选中要删除的分级")
        if len(self._stages) <= 3: return QMessageBox.warning(self, "无法删除", "至少保留 3 个分级。")
        r = sel[0].row()
        if 0 <= r < len(self._stages): self._stages.pop(r); self._rebuild_stage_table()

    def _on_stage_selection_changed(self): self._update_compound_ui()

    def _update_compound_ui(self):
        for c in self._compound_cards:
            self._cl.removeWidget(c); c.deleteLater()
        self._compound_cards.clear()
        sel = self.st.selectedItems(); idx = sel[0].row() if sel else None
        if idx is None: self._sel_lbl.setText("未选中任何分级（请点击表格中[序号]或[化合物]列选中行）"); return
        s = self._stages[idx]; self._sel_lbl.setText(f"当前分级：{s['name']}（{s['low']:.2f} – {s['high']:.2f} μm）")
        elem = self.ec.currentText(); nucs = _element_nuclides_map.get(elem, [])
        # 气溶胶选项
        opts = set()
        for n in nucs:
            df = get_nuclide_df(n)
            if df is not None and 'aerosol_type' in df.columns: opts.update(df['aerosol_type'].dropna().unique())
        os2 = sorted(opts)
        for cd in s['compounds']:
            card = CompactCompoundCard()
            card.update_aerosol_options(os2)
            card.update_nuclide_options(nucs)          # ★ 核素下拉读取自数据文件
            card.set_data(cd)
            # 设置变更回调（父级在核素增删后同步数据）
            card._on_changed_cb = lambda si=idx: self._on_cd(si)
            # 字段变更信号
            card.ce.textChanged.connect(lambda t, si=idx: self._on_cd(si))
            card.af.valueChanged.connect(lambda v, si=idx: self._on_cd(si))
            card.ac.currentTextChanged.connect(lambda t, si=idx: self._on_cd(si))
            db = QPushButton("✕ 删除化合物"); db.setFixedHeight(22)
            db.setStyleSheet("color:red; font-size:9px; border:1px solid #e0c0c0; border-radius:3px;")
            db.clicked.connect(lambda ch, c=card: self._rm_card(c))
            pos = self._cl.count()-1; self._cl.insertWidget(pos, card); self._compound_cards.append(card)
            self._cl.insertWidget(pos+1, db); self._compound_cards.append(db)

    def _on_cd(self, si):
        if si is None or si >= len(self._stages): return
        comps = [w.get_data() for w in self._compound_cards if isinstance(w, CompactCompoundCard)]
        if comps: self._stages[si]['compounds'] = comps
        if si < self.st.rowCount():
            self.st.item(si, 4).setText(str(len(comps)))

    def _rm_card(self, card):
        si = self.st.selectedItems()
        if not si: return; si = si[0].row()
        if len(self._stages[si]['compounds']) <= 1: return QMessageBox.information(self, "提示", "至少保留一个化合物")
        ci = next((i for i, w in enumerate(self._compound_cards) if w is card), None)
        if ci is None: return
        nb = self._compound_cards[ci+1] if ci+1 < len(self._compound_cards) else None
        self._cl.removeWidget(card); card.deleteLater()
        if nb and not isinstance(nb, CompactCompoundCard):
            self._cl.removeWidget(nb); nb.deleteLater(); self._compound_cards.remove(nb)
        self._compound_cards.remove(card); self._on_cd(si)

    def _add_compound_to_current_stage(self):
        si = self.st.selectedItems()
        if not si: return QMessageBox.information(self, "提示", "请先选中分级")
        si = si[0].row(); self._stages[si]['compounds'].append(dict(DEFAULT_COMP))
        self._update_compound_ui()

    def _copy_compounds_to_all(self):
        si = self.st.selectedItems()
        if not si: return QMessageBox.information(self, "提示", "请先选中源分级")
        si = si[0].row(); src = [dict(c) for c in self._stages[si]['compounds']]
        for s in self._stages: s['compounds'] = [dict(c) for c in src]
        self._rebuild_stage_table()
        QMessageBox.information(self, "完成", f"已复制到全部 {len(self._stages)} 个分级")

    # ── 数据扫描 ──
    def _scan_data(self):
        self.sb.showMessage("正在扫描核素数据文件…"); scan_nuclides()
        elems = sorted(_element_nuclides_map.keys()); self.ec.clear(); self.ec.addItems(elems)
        if elems: self._on_elem_changed(elems[0])
        self.sb.showMessage(f"已加载 {len(elems)} 个元素  |  就绪")

    def _on_input_mode_changed(self):
        man = self.rb_manual.isChecked(); self.fr.setVisible(not man)
        if man: self._stages = []; self._init_default_stages(); self._rebuild_stage_table()
        self._lock_stage_inputs(not man)

    def _lock_stage_inputs(self, locked):
        """锁定/解锁分级表格和化合物编辑（文件导入时锁定）"""
        self.st.setEnabled(not locked)
        self.btn_add_s.setEnabled(not locked)
        self.btn_del_s.setEnabled(not locked)
        self.btn_cp.setEnabled(not locked)
        self.ec.setEnabled(not locked)
        self._cc.setEnabled(not locked)
        for card in self._compound_cards:
            card.setEnabled(not locked)
        for i in range(self.st.rowCount()):
            for j in (1, 2, 3):
                w = self.st.cellWidget(i, j)
                if w:
                    w.setReadOnly(locked)
                    if locked:
                        w.setStyleSheet("QDoubleSpinBox { background:#f0f0f0; border:1px solid #ddd; }")
                    else:
                        w.setStyleSheet("")

    def _browse_file(self):
        p, _ = QFileDialog.getOpenFileName(self, "选择数据文件", "", "数据文件 (*.csv *.xlsx *.xls)")
        if not p: return
        self.fp_edit.setText(p)
        try:
            self._file_df = pd.read_csv(p) if p.endswith('.csv') else pd.read_excel(p)
            ws = sorted(self._file_df['Workshop'].unique())
            self.ws_cb.blockSignals(True); self.ws_cb.clear(); self.ws_cb.addItems([str(w) for w in ws]); self.ws_cb.blockSignals(False)
            if ws: self._on_ws_changed(str(ws[0]))
        except Exception as e: QMessageBox.warning(self, "读取失败", str(e))

    def _on_ws_changed(self, ws):
        if self._file_df is None: return
        dfw = self._file_df[self._file_df['Workshop'].astype(str) == ws]
        sids = sorted(dfw['SamplingID'].unique())
        self.sid_cb.blockSignals(True); self.sid_cb.clear(); self.sid_cb.addItems([str(s) for s in sids]); self.sid_cb.blockSignals(False)
        if sids: self._on_sid_changed(str(sids[0]))

    def _on_sid_changed(self, sid):
        if self._file_df is None: return
        ws = self.ws_cb.currentText()
        m = (self._file_df['Workshop'].astype(str) == ws) & (self._file_df['SamplingID'].astype(str) == sid)
        rows = self._file_df[m]
        if rows.empty: return
        row = rows.iloc[0]
        for i, col in enumerate(CONC_COLS):
            if col in row.index and i < len(self._stages):
                try: self._stages[i]['conc'] = float(row[col])
                except: self._stages[i]['conc'] = 0.0
        try:
            cs = str(row.get('compounds', '')); fs = str(row.get('activity_fractions', '')); ns = str(row.get('nuclide_abundances', ''))
            if cs and cs.lower() != 'nan':
                cl = cs.split(';'); fl = fs.split(';'); nl = ns.split(';')
                fcs = []
                for ci in range(len(cl)):
                    cn = cl[ci].strip()
                    if not cn: continue
                    fv = float(fl[ci]) if ci < len(fl) else 1.0
                    ns2 = nl[ci].strip() if ci < len(nl) else "U_238:0.993|U_235:0.007"
                    r = _COMPOUND_SHORT_TO_FULL.get(cn); aero = r[1] if r else 'Intermediate Type M/S'
                    nucs = {}
                    for pp in ns2.split('|'):
                        pp = pp.strip()
                        if ':' in pp:
                            n2, v2 = pp.split(':', 1)
                            try: nucs[n2.strip()] = float(v2.strip())
                            except: pass
                    if not nucs: nucs = {'U_238': 0.993, 'U_235': 0.007}
                    fcs.append({'compound': cn, 'aerosol_type': aero, 'activity_fraction': fv, 'nuclides': nucs})
                if fcs:
                    for s in self._stages: s['compounds'] = [dict(c) for c in fcs]
        except: pass
        self._rebuild_stage_table()
        if not self.rb_manual.isChecked():
            self._lock_stage_inputs(True)

    def _on_elem_changed(self, elem):
        nucs = _element_nuclides_map.get(elem, []); opts = set()
        for n in nucs:
            df = get_nuclide_df(n)
            if df is not None and 'aerosol_type' in df.columns: opts.update(df['aerosol_type'].dropna().unique())
        os2 = sorted(opts)
        for c in self._compound_cards:
            if isinstance(c, CompactCompoundCard):
                c.update_aerosol_options(os2)
                c.update_nuclide_options(nucs)
        self._current_elem = elem

    def _update_mask_info(self):
        mt = self.mc.currentText(); p = MASK_PARAMS.get(mt, {})
        self.mil.setText(f"过滤效率 {p.get('filtration_efficiency',0):.0%}  泄漏率 {p.get('leakage_rate',1):.0%}")

    def _update_method_hints(self):
        for rb, hl in self._hl.items():
            hl.setStyleSheet("color:#1a5276; font-size:10px; font-weight:bold; margin-left:16px;" if rb.isChecked() else "color:#999; font-size:10px; margin-left:16px;")
        self._mpw.setVisible(self.rb_mod.isChecked())

    # ── 拟合 ──
    def _on_fit(self):
        l, h, m, c, rc = self._get_stage_arrays()
        if np.sum(rc) == 0: return QMessageBox.warning(self, "数据为空", "请先输入浓度数据。")
        if len(self._stages) < 3: return QMessageBox.warning(self, "分级不足", "至少需要 3 个分级。")
        try:
            ec = _apply_efficiency_dynamic(l, h, rc); tc = float(np.sum(ec))
            old = _patch_xf_stages(l, h, c)
            q = data_quality_assessment(rc)
            au, gu, tu = fit_unimodal(ec)
            a1, g1, f1, a2, g2, tb = fit_bimodal(ec)
            D50, GSD, _, _, R2 = fit_linear_probit(rc)
            aicu, bicu = calc_unimodal_stats(au, gu, tu, ec)
            aicb, bicb = (np.nan, np.nan)
            if not np.isnan(a2): aicb, bicb = calc_bimodal_stats(a1, g1, f1, a2, g2, tb, ec)
            Dp, Gp, Rp, ap, bp, npb = calc_probit_stats(rc)
            rt = a1/a2 if not np.isnan(a2) and a2 > 0 else np.nan
            rec = recommend_method(a2, rt, f1, R2, aic_uni=aicu, bic_uni=bicu, aic_bi=aicb, bic_bi=bicb,
                                   quality_flag=q['quality_flag'], n_nonzero=q['n_nonzero'])
            _restore_xf_stages(old)
            self._fit_auto = {'raw': rc, 'ec': ec, 'tc': tc, 'au': au, 'gu': gu, 'tu': tu,
                              'a1': a1, 'g1': g1, 'f1': f1, 'a2': a2, 'g2': g2, 'D50': D50, 'GSD': GSD, 'R2': R2,
                              'rec': rec, 'aicu': aicu, 'bicu': bicu, 'aicb': aicb, 'bicb': bicb, 'ap': ap, 'q': q}
            self.pc.plot_fitting(rc, ec, au, gu, tu, a1, g1, f1, a2, g2, tb, tc, D50, GSD, R2, False, l, h, m, c)
            ns = len(self._stages)
            ps = [f"【推荐方法】{rec}", f"单峰 AMAD={au:.3f} μm  GSD={gu:.3f}  AIC={aicu:.1f}  BIC={bicu:.1f}"]
            if not np.isnan(a2): ps.append(f"双峰：粗峰={a1:.3f}  细峰={a2:.3f}  粗峰={f1:.2%}  AIC={aicb:.1f}  BIC={bicb:.1f}")
            if not np.isnan(D50): ps.append(f"正态概率图 AMAD={D50:.3f} μm  GSD={GSD:.3f}  R²={R2:.4f}")
            ps.append(f"数据质量: {q['quality_flag']}（非零级数={q['n_nonzero']}/{ns}）")
            self.fil.setText("  \n".join(ps))
            self.le.append(f"[拟合] 单峰 AMAD={au:.3f} AIC={aicu:.1f} | 双峰粗峰={a1:.3f} AIC={aicb:.1f} | 数据质量={q['quality_flag']}")
            self.sb.showMessage(f"✓ 拟合完成 | 推荐: {rec}", 5000)
        except Exception as e:
            import traceback; self.le.append(f"[拟合失败] {traceback.format_exc()}"); QMessageBox.warning(self, "拟合失败", str(e))

    # ── 核心计算入口 ──
    def _on_calc(self):
        l, h, m, c, rc = self._get_stage_arrays()
        if np.sum(rc) == 0: return QMessageBox.warning(self, "数据为空", "请先输入浓度数据。")
        if len(self._stages) < 3: return QMessageBox.warning(self, "分级不足", "至少需要 3 个分级。")
        self.cb.setEnabled(False); self.sb.showMessage("计算中…"); self.le.clear()
        mt = self.mc.currentText(); br = self.brs.value(); wh = self.ws_spin.value()
        elem = self.ec.currentText(); ne = _element_nuclides_map.get(elem, [])
        method = ('std' if self.rb_std.isChecked() else 'modal' if self.rb_mod.isChecked() else
                  'probit' if self.rb_pro.isChecked() else 'stage')
        up = 'bimodal' if self.rb_pb.isChecked() else 'unimodal'
        snap = [{k: (v if k != 'compounds' else [dict(c) for c in v]) for k, v in s.items()} for s in self._stages]
        self._calc_thread = CalcThread(self._do_calc, snap, mt, br, wh, method, ne, up)
        self._calc_thread.result_ready.connect(self._on_result); self._calc_thread.error_signal.connect(self._on_error)
        self._calc_thread.start()

    # ═══════════════ 核心计算（线程内运行）═══════════════
    @staticmethod
    def _locate_stage_dose(detail_ls):
        """在 detail 列表中定位最高剂量记录"""
        if not detail_ls: return 0.0
        return max(float(d.get('dose_sv', 0) or 0) for d in detail_ls)

    @staticmethod
    def _do_calc(snap, mask_type, br, wh, method, nuclide_names, peak_mode):
        stages = snap
        n = len(stages)
        low = np.array([s['low'] for s in stages], dtype=float)
        high = np.array([s['high'] for s in stages], dtype=float)
        mid = np.sqrt(low * high)
        cut = high[:-1].copy() if n > 1 else np.array([])
        raw_concs = np.array([s['conc'] for s in stages], dtype=float)

        # 效率校正
        eff_arr = _build_efficiency(low, high)
        corrected = raw_concs / eff_arr
        total_corrected = float(np.sum(corrected))

        # --- 拟合（仅 modal 和 probit 需要）---
        fit_info = {}
        amad_for_stages = np.full(n, np.nan)  # 每级使用的 AMAD
        gsd_for_stages = np.full(n, np.nan)
        coarse_frac_for_stages = np.full(n, np.nan)  # 双峰粗峰权重

        if method == 'std':
            amad_for_stages[:] = 5.0
            gsd_for_stages[:] = 2.5
            fit_info['desc'] = '国标单AMAD法（固定 5 μm）'
        elif method == 'modal':
            old = _patch_xf_stages(low, high, cut)
            try:
                au, gu, tu = fit_unimodal(corrected)
                a1, g1, f1, a2, g2, tb = fit_bimodal(corrected)
            finally:
                _restore_xf_stages(old)
            fit_info['au'] = au; fit_info['gu'] = gu; fit_info['tu'] = tu
            fit_info['a1'] = a1; fit_info['g1'] = g1; fit_info['f1'] = f1
            fit_info['a2'] = a2; fit_info['g2'] = g2; fit_info['tb'] = tb

            if peak_mode == 'unimodal' or np.isnan(a2):
                amad_for_stages[:] = au; gsd_for_stages[:] = gu
                fit_info['desc'] = f'多模态-单峰 AMAD={au:.3f} μm'
            else:
                # 双峰：逐级 PDF 权重分配
                fit_info['desc'] = f'多模态-双峰 粗峰={a1:.3f} 细峰={a2:.3f} μm'
                for i in range(n):
                    d_i = mid[i]
                    pdf1 = lognormal_pdf(d_i, a1, g1)
                    pdf2 = lognormal_pdf(d_i, a2, g2)
                    denom = f1 * pdf1 + (1.0 - f1) * pdf2
                    if denom > 1e-30:
                        w1 = f1 * pdf1 / denom
                    else:
                        w1 = 0.5
                    amad_for_stages[i] = a1  # 粗峰 AMAD
                    gsd_for_stages[i] = g1
                    coarse_frac_for_stages[i] = w1  # 粗峰权重
        elif method == 'probit':
            old = _patch_xf_stages(low, high, cut)
            try:
                D50, GSD, D84, D16, R2 = fit_linear_probit(raw_concs)
            finally:
                _restore_xf_stages(old)
            if np.isnan(D50):
                D50 = 5.0; GSD = 2.5
            amad_for_stages[:] = D50; gsd_for_stages[:] = GSD
            fit_info['D50'] = D50; fit_info['GSD'] = GSD; fit_info['R2'] = R2
            fit_info['desc'] = f'正态概率图法 AMAD={D50:.3f} μm'
        elif method == 'stage':
            amad_for_stages[:] = mid  # 各用各的中值粒径
            gsd_for_stages[:] = 2.0
            fit_info['desc'] = '逐级独立法（各级独立均质源）'

        # --- 逐级 · 逐化合物 · 逐核素计算剂量 ---
        details = []
        total_dose = 0.0

        for i, s in enumerate(stages):
            d_i = float(mid[i])
            conc_i = float(corrected[i])
            pf_i = calc_mask_pf(mask_type, d_i)
            compounds = s.get('compounds', [])

            for comp in compounds:
                cname = comp.get('compound', '')
                aero = comp.get('aerosol_type', '')
                af = float(comp.get('activity_fraction', 1.0))
                if af <= 0:
                    continue
                nuclides = comp.get('nuclides', {})
                if not nuclides:
                    continue

                for nuc_name, abundance in nuclides.items():
                    ab = float(abundance)
                    if ab <= 0:
                        continue
                    nuc_name_norm = _normalize_nuclide(nuc_name)
                    df_nuc = get_nuclide_df(nuc_name_norm, inhalation_only=True)

                    contribute_conc = conc_i * af * ab

                    # 双峰模式：逐级权重分配
                    if (method == 'modal' and peak_mode == 'bimodal'
                            and not np.isnan(fit_info.get('a2', np.nan))):
                        w1 = float(coarse_frac_for_stages[i])
                        w2 = 1.0 - w1
                        e1 = interp_dose_coeff(df_nuc, aero, float(fit_info['a1']))
                        e2 = interp_dose_coeff(df_nuc, aero, float(fit_info['a2']))
                        if e1 is None and e2 is None:
                            continue
                        e1 = e1 or 0.0; e2 = e2 or 0.0
                        dose_sv_this = contribute_conc * (w1 * e1 + w2 * e2) * br * wh * pf_i
                        amad_label = f"粗峰{fit_info['a1']:.2f}(w={w1:.3f})+细峰{fit_info['a2']:.2f}(w={w2:.3f})"
                        e_display = (e1 + e2) / 2.0
                    else:
                        amad_val = float(amad_for_stages[i])
                        e_val = interp_dose_coeff(df_nuc, aero, amad_val)
                        if e_val is None:
                            continue
                        dose_sv_this = contribute_conc * e_val * br * wh * pf_i
                        amad_label = f"{amad_val:.2f}"
                        e_display = e_val

                    total_dose += dose_sv_this
                    details.append({
                        'stage': s['name'], 'd_um': d_i,
                        'compound': cname, 'aerosol_type': aero,
                        'nuclide': nuc_name_norm, 'amad_used': amad_label,
                        'e_sv_per_bq': e_display,
                        'conc_contrib': contribute_conc,
                        'dose_sv': dose_sv_this, 'pf': pf_i,
                    })

        return {
            'method': method, 'peak_mode': peak_mode, 'mask_type': mask_type,
            'br': br, 'wh': wh,
            'raw_concs': raw_concs.tolist(), 'corrected': corrected.tolist(),
            'total_corrected': total_corrected, 'total_dose': total_dose,
            'details': details, 'fit_info': fit_info,
            'n_stages': n,
            'amad_for_stages': amad_for_stages.tolist(),
            'gsd_for_stages': gsd_for_stages.tolist(),
        }

    # ── 结果展示 ──
    def _on_result(self, result):
        self.cb.setEnabled(True); self.sb.showMessage("计算完成", 5000)
        self._fit_result = result
        self.le.append(f"=== {datetime.datetime.now().strftime('%H:%M:%S')} 计算完成 ===")
        self.le.append(f"方法: {result['fit_info'].get('desc','')}")
        self.le.append(f"总剂量: {result['total_dose']:.6e} Sv")
        self.le.append(f"详情记录数: {len(result['details'])}")
        self.le.append("")

        # 更新剂量详情表格 Tab2
        self.rt2.setRowCount(len(result['details']))
        for ri, d in enumerate(result['details']):
            for ci, key in enumerate(['compound', 'aerosol_type', 'nuclide',
                                       'amad_used', 'e_sv_per_bq', 'conc_contrib', 'dose_sv']):
                val = d.get(key, '')
                if key == 'e_sv_per_bq' and isinstance(val, (int, float)):
                    val = f"{val:.4e}"
                elif key == 'conc_contrib' and isinstance(val, (int, float)):
                    val = f"{val:.4e}"
                elif key == 'dose_sv' and isinstance(val, (int, float)):
                    val = f"{val:.4e}"
                self.rt2.setItem(ri, ci, QTableWidgetItem(str(val)))

        # 更新摘要标签
        info = result['fit_info']
        lines = [
            f"══════ 计算结果摘要 ══════",
            f"方法: {info.get('desc','')}",
            f"总有效浓度: {result['total_corrected']:.4f} Bq/m³",
            f"总剂量: {result['total_dose']:.6e} Sv",
            f"分级数: {result['n_stages']}",
            f"呼吸速率: {result['br']:.2f} m³/h  |  工作时长: {result['wh']:.1f} h",
            f"口罩: {result['mask_type']}",
        ]
        self.sl.setText("\n".join(lines))

        # 更新拟合图（仅在 modal 或 probit 方法时）
        if result['method'] in ('modal', 'probit'):
            self._redraw_fit_plot(result)

    def _redraw_fit_plot(self, result):
        """从 calc 结果重绘拟合图"""
        try:
            l, h, m, c, _ = self._get_stage_arrays()
            raw = np.array(result['raw_concs'])
            corr = np.array(result['corrected'])
            tc = result['total_corrected']
            fi = result['fit_info']
            au = fi.get('au', 5.0); gu = fi.get('gu', 2.5); tu = fi.get('tu', tc)
            a1 = fi.get('a1', np.nan); g1 = fi.get('g1', np.nan)
            f1v = fi.get('f1', np.nan); a2 = fi.get('a2', np.nan)
            g2 = fi.get('g2', np.nan); tb = fi.get('tb', tc)
            D50 = fi.get('D50', np.nan); GSD = fi.get('GSD', np.nan)
            R2 = fi.get('R2', np.nan)
            self.pc.plot_fitting(raw, corr, au, gu, tu, a1, g1, f1v, a2, g2,
                                 tb, tc, D50, GSD, R2, False, l, h, m, c)
            self.fil.setText(f"拟合参数: {fi.get('desc','')}")
        except Exception as e:
            self.le.append(f"[重绘拟合图失败] {e}")

    def _on_error(self, err):
        self.cb.setEnabled(True); self.sb.showMessage("计算出错", 5000)
        self.le.append(f"[错误] {err}")
        QMessageBox.critical(self, "计算错误", err)

    def _clear_results(self):
        self._fit_result = None; self.sl.setText("剂量结果将在计算后显示")
        self.rt2.setRowCount(0); self.pc.fig.clear(); self.pc.draw()
        self.fil.setText("拟合参数将在计算后显示"); self.le.clear()
        self.sb.showMessage("已清空结果")

    # ── 累计剂量表格 ──
    def _add_to_accum(self):
        if self._fit_result is None:
            return QMessageBox.information(self, "提示", "请先完成计算。")
        rr = self._fit_result
        fi = rr['fit_info']
        method_desc = fi.get('desc', '')
        # 取最高剂量行作为摘要
        top = max(rr['details'], key=lambda d: d.get('dose_sv', 0)) if rr['details'] else {}
        cpd = top.get('compound', ''); nuc = top.get('nuclide', '')
        label = f"{cpd}/{nuc}"
        amad = top.get('amad_used', '')
        pf0 = top.get('pf', 1.0)

        entry = {
            'method': method_desc, 'compound_label': label,
            'amad': str(amad), 'pf': pf0, 'dose_sv': rr['total_dose'],
            'details': rr['details'], 'fit_info': fi,
        }
        self._calc_history.append(entry)
        self._refresh_accum_table()

    def _del_accum_row(self):
        sel = self.at.selectedItems()
        if not sel:
            return QMessageBox.information(self, "提示", "请先选中要删除的行。")
        r = sel[0].row()
        if 0 <= r < len(self._calc_history):
            self._calc_history.pop(r)
            self._refresh_accum_table()

    def _clear_accum(self):
        if not self._calc_history:
            return
        if QMessageBox.question(self, "确认", "确定清空所有累计记录？") == QMessageBox.Yes:
            self._calc_history.clear()
            self._refresh_accum_table()

    def _refresh_accum_table(self):
        self.at.setRowCount(len(self._calc_history))
        total = 0.0
        for ri, e in enumerate(self._calc_history):
            for ci, key in enumerate(['method', 'compound_label', 'amad', 'pf', 'dose_sv']):
                val = e.get(key, '')
                if key == 'dose_sv' and isinstance(val, (int, float)):
                    total += val; val = f"{val:.4e}"
                elif key == 'pf' and isinstance(val, (int, float)):
                    val = f"{val:.3f}"
                self.at.setItem(ri, ci, QTableWidgetItem(str(val)))
        self.atl.setText(f"累计 {len(self._calc_history)} 条  |  合计剂量: {total:.6e} Sv")

    # ── 导出 ──
    def _check_result(self):
        if self._fit_result is None and not self._calc_history:
            QMessageBox.information(self, "提示", "请先完成计算或加入累计。"); return False
        return True

    def _build_export_data(self, include_all_history=True):
        """构建导出用的结构化数据"""
        rows = []
        if include_all_history:
            for i, e in enumerate(self._calc_history):
                for d in e.get('details', []):
                    rows.append({**d, 'record_index': i + 1,
                                 'method': e.get('method', ''),
                                 'total_dose': e.get('dose_sv', 0)})
        elif self._fit_result:
            for d in self._fit_result.get('details', []):
                rows.append({**d, 'record_index': 1,
                             'method': self._fit_result.get('fit_info', {}).get('desc', ''),
                             'total_dose': self._fit_result.get('total_dose', 0)})
        return rows

    def _do_export_xlsx(self, path):
        import openpyxl
        try:
            wb = openpyxl.Workbook()
            ws = wb.active; ws.title = "剂量详情"
            headers = ['序号', '方法', '分级', '粒径(μm)', '化合物', '气溶胶',
                       '核素', 'AMAD', 'e(Sv/Bq)', '活度贡献', '剂量(Sv)', '总剂量(Sv)']
            for c, h in enumerate(headers, 1):
                ws.cell(row=1, column=c, value=h)
            rows = self._build_export_data(True)
            for ri, r in enumerate(rows, 2):
                vals = [ri - 1, r.get('method', ''), r.get('stage', ''),
                        r.get('d_um', ''), r.get('compound', ''), r.get('aerosol_type', ''),
                        r.get('nuclide', ''), r.get('amad_used', ''),
                        r.get('e_sv_per_bq', ''), r.get('conc_contrib', ''),
                        r.get('dose_sv', ''), r.get('total_dose', '')]
                for c, v in enumerate(vals, 1):
                    ws.cell(row=ri, column=c, value=v)
            wb.save(path)
            return True, path
        except Exception as e:
            return False, str(e)

    def _do_export_csv(self, path):
        import csv
        rows = self._build_export_data(True)
        if not rows:
            return False, "无数据可导出"
        cols = ['stage', 'd_um', 'compound', 'aerosol_type', 'nuclide',
                'amad_used', 'e_sv_per_bq', 'conc_contrib', 'dose_sv', 'total_dose']
        try:
            with open(path, 'w', newline='', encoding='utf-8-sig') as f:
                w = csv.DictWriter(f, fieldnames=cols, extrasaction='ignore')
                w.writeheader(); w.writerows(rows)
            return True, path
        except Exception as e:
            return False, str(e)

    def _do_export_log(self, path):
        try:
            with open(path, 'w', encoding='utf-8') as f:
                f.write(f"剂量计算日志  {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
                f.write("=" * 60 + "\n\n")
                if self._fit_result:
                    rr = self._fit_result
                    f.write(f"方法: {rr['fit_info'].get('desc','')}\n")
                    f.write(f"总剂量: {rr['total_dose']:.6e} Sv\n")
                    f.write(f"呼吸速率: {rr['br']:.2f} m³/h  工作时长: {rr['wh']:.1f} h\n")
                    f.write(f"口罩: {rr['mask_type']}\n\n")
                f.write("--- 累计记录 ---\n")
                for i, e in enumerate(self._calc_history):
                    f.write(f"[{i+1}] {e['method']} | {e['compound_label']} | {e['dose_sv']:.6e} Sv\n")
                f.write("\n--- 原始日志 ---\n")
                f.write(self.le.toPlainText() if hasattr(self, 'le') else '')
            return True, path
        except Exception as e:
            return False, str(e)

    def _export_xlsx(self):
        if not self._check_result(): return
        p, _ = QFileDialog.getSaveFileName(self, "导出 Excel", "dose_report.xlsx",
                                            "Excel (*.xlsx)")
        if not p: return
        ok, msg = self._do_export_xlsx(p)
        if ok:
            self.le.append(f"[导出] Excel → {p}")
            self.sb.showMessage(f"已导出: {os.path.basename(p)}")
            QMessageBox.information(self, "导出完成", f"已保存到:\n{p}")
        else:
            QMessageBox.warning(self, "导出失败", msg)

    def _export_csv(self):
        if not self._check_result(): return
        p, _ = QFileDialog.getSaveFileName(self, "导出 CSV", "dose_report.csv",
                                            "CSV (*.csv)")
        if not p: return
        ok, msg = self._do_export_csv(p)
        if ok:
            self.le.append(f"[导出] CSV → {p}")
            self.sb.showMessage(f"已导出: {os.path.basename(p)}")
            QMessageBox.information(self, "导出完成", f"已保存到:\n{p}")
        else:
            QMessageBox.warning(self, "导出失败", msg)

    def _export_log(self):
        if not self._check_result(): return
        p, _ = QFileDialog.getSaveFileName(self, "导出日志", "dose_log.txt",
                                            "文本 (*.txt)")
        if not p: return
        ok, msg = self._do_export_log(p)
        if ok:
            self.le.append(f"[导出] 日志 → {p}")
            self.sb.showMessage(f"已导出: {os.path.basename(p)}")
            QMessageBox.information(self, "导出完成", f"已保存到:\n{p}")
        else:
            QMessageBox.warning(self, "导出失败", msg)

    def _export_all(self):
        if not self._check_result(): return
        from pathlib import Path as P
        d = QFileDialog.getExistingDirectory(self, "选择导出目录")
        if not d: return
        base = P(d)
        t = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
        results = []
        for suf, fnc in [('.xlsx', self._do_export_xlsx), ('.csv', self._do_export_csv), ('.txt', self._do_export_log)]:
            fp = str(base / f"dose_report_{t}{suf}")
            ok, msg = fnc(fp)
            results.append((suf, ok, msg))
        # 保存拟合图
        try:
            img_p = str(base / f"fitting_plot_{t}.png")
            self.pc.fig.savefig(img_p, dpi=150)
            results.append(('.png', True, img_p))
        except Exception as e:
            results.append(('.png', False, str(e)))
        msg = "导出结果:\n" + "\n".join([f"  {s}: {'✓' if ok else '✗'} {m}" for s, ok, m in results])
        self.le.append(f"[导出全部] → {d}")
        self.sb.showMessage(f"已全部导出到: {d}")
        QMessageBox.information(self, "导出完成", msg)


# ==================== 入口 ====================
def main():
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    QApplication.setAttribute(Qt.AA_UseHighDpiPixmaps, True)
    app = QApplication(sys.argv)
    app.setStyle('Fusion')
    palette = QPalette()
    palette.setColor(QPalette.Window, QColor(245, 248, 252))
    palette.setColor(QPalette.Base, QColor(255, 255, 255))
    app.setPalette(palette)
    window = DoseCalcApp()
    window.show()
    sys.exit(app.exec_())

if __name__ == '__main__':
    main()

